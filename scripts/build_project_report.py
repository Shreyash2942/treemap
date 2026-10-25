"""Build the expanded APA-style project report as PDF and editable Word files."""

from __future__ import annotations

import argparse
import json
import re
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Flowable,
    Frame,
    Image,
    LongTable,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
WIDTH = 468


def blocks(source: str):
    """Read the report's headings, paragraphs, tables, images, and page markers."""
    lines = source.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line.startswith("<!--"):
            yield "marker", line.removeprefix("<!--").removesuffix("-->").strip()
        elif match := re.fullmatch(r"(#{1,3}) (.+)", line):
            yield "heading", (len(match[1]), match[2])
        elif match := re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line):
            yield "image", (match[1], match[2])
        elif line.startswith("|"):
            rows = [line]
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(lines[index].strip())
                index += 1
            cells = [[cell.strip() for cell in row.strip("|").split("|")] for row in rows]
            yield (
                "table",
                [row for row in cells if not all(re.fullmatch(r":?-+:?", c) for c in row)],
            )
        else:
            paragraph = [line]
            while index < len(lines) and lines[index].strip():
                paragraph.append(lines[index].strip())
                index += 1
            yield "paragraph", " ".join(paragraph)


def spans(text: str):
    """Yield plain, bold, italic, and inline-code spans without interpreting HTML."""
    for part in re.split(r"(\*\*.*?\*\*|\*[^*]+\*|`[^`]+`)", text):
        if part.startswith("**") and part.endswith("**"):
            yield part[2:-2], "bold"
        elif part.startswith("*") and part.endswith("*"):
            yield part[1:-1], "italic"
        elif part.startswith("`") and part.endswith("`"):
            yield part[1:-1], "code"
        elif part:
            yield part, "plain"


def markup(text: str) -> str:
    result = []
    for value, kind in spans(text):
        value = escape(value)
        if kind == "bold":
            value = f"<b>{value}</b>"
        elif kind == "italic":
            value = f"<i>{value}</i>"
        elif kind == "code":
            value = f'<font name="Courier" size="10">{value}</font>'
        result.append(value)
    return re.sub(
        r"https://[^\s<]+",
        lambda match: f'<link href="{match[0]}" color="black">{match[0]}</link>',
        "".join(result),
    )


def register_font() -> str:
    folder = Path("C:/Windows/Fonts")
    faces = {
        "ReportTimes": "times.ttf",
        "ReportTimes-Bold": "timesbd.ttf",
        "ReportTimes-Italic": "timesi.ttf",
        "ReportTimes-BoldItalic": "timesbi.ttf",
    }
    if not all((folder / name).is_file() for name in faces.values()):
        return "Times-Roman"
    for name, filename in faces.items():
        pdfmetrics.registerFont(TTFont(name, str(folder / filename)))
    pdfmetrics.registerFontFamily(
        "ReportTimes",
        normal="ReportTimes",
        bold="ReportTimes-Bold",
        italic="ReportTimes-Italic",
        boldItalic="ReportTimes-BoldItalic",
    )
    return "ReportTimes"


class Marker(Flowable):
    """Record pagination without inserting visible content or padding pages."""

    def __init__(self, name: str):
        super().__init__()
        self.name = name

    def draw(self):
        pass


class ReportDocument(BaseDocTemplate):
    def __init__(self, path: Path, title: str, font: str):
        super().__init__(
            str(path),
            pagesize=letter,
            leftMargin=72,
            rightMargin=72,
            topMargin=72,
            bottomMargin=72,
            title=title,
            author="",
            invariant=1,
        )
        self.markers = {}
        frame = Frame(
            72, 72, WIDTH, 648, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0
        )

        def header(canvas, document):
            canvas.setFont(font, 12)
            canvas.drawRightString(540, 756, str(document.page))

        self.addPageTemplates(PageTemplate(id="APA", frames=[frame], onPage=header))

    def afterFlowable(self, flowable):
        if isinstance(flowable, Marker):
            self.markers[flowable.name] = self.page


def column_weights(rows):
    if len(rows[0]) == 5:
        return [1.8, 1.175, 1.175, 1.175, 1.175]
    if rows[0][0] == "ID":
        return [0.65, 1.6, 4.25]
    return [1.8, 0.9, 3.8]


def add_word_text(paragraph, text):
    for value, kind in spans(text):
        run = paragraph.add_run(value)
        run.bold = kind == "bold"
        run.italic = kind == "italic"
        if kind == "code":
            run.font.name = "Courier New"
            run.font.size = Pt(10)


def word_table(document, rows):
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.autofit = False
    for column, width in zip(table.columns, column_weights(rows), strict=True):
        column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single" if edge in {"top", "bottom"} else "nil")
        border.set(qn("w:sz"), "6")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, width, value in zip(row.cells, column_weights(rows), values, strict=True):
            cell.width = Inches(width)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_with_next = index == 0
            add_word_text(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(10)
                if index == 0:
                    run.bold = True
            if index == 0:
                cell_borders = OxmlElement("w:tcBorders")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                cell_borders.append(bottom)
                cell._tc.get_or_add_tcPr().append(cell_borders)


def build_report(source: Path, output: Path) -> dict:
    text = source.read_text(encoding="utf-8")
    parsed = list(blocks(text))
    title = parsed[0][1][1]
    output.mkdir(parents=True, exist_ok=True)
    font = register_font()
    bold = "ReportTimes-Bold" if font == "ReportTimes" else "Times-Bold"
    body = ParagraphStyle(
        "Body",
        fontName=font,
        fontSize=12,
        leading=24,
        firstLineIndent=36,
        alignment=TA_LEFT,
        spaceBefore=0,
        spaceAfter=0,
        allowWidows=0,
        allowOrphans=0,
    )
    no_indent = ParagraphStyle("NoIndent", parent=body, firstLineIndent=0)
    heading = ParagraphStyle(
        "Heading",
        parent=no_indent,
        fontName=bold,
        alignment=TA_CENTER,
        keepWithNext=True,
    )
    left_heading = ParagraphStyle("Subheading", parent=heading, alignment=TA_LEFT)
    reference = ParagraphStyle("Reference", parent=body, leftIndent=36, firstLineIndent=-36)
    caption_style = ParagraphStyle("Caption", parent=no_indent, keepWithNext=True)
    cell = ParagraphStyle("Cell", parent=no_indent, fontSize=10, leading=13)

    word = Document()
    word.core_properties.title = title
    word.core_properties.author = ""
    word.core_properties.last_modified_by = ""
    word.core_properties.comments = ""
    section = word.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    for name in ("Normal", "Heading 1", "Heading 2", "Title"):
        style = word.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = None
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.widow_control = True
    word.styles["Normal"].paragraph_format.first_line_indent = Inches(0.5)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent = Inches(0)
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    header._p.append(page_field)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    word.settings.element.append(update)

    story = []
    state = "abstract"
    caption_pending = False

    def add_paragraph(value, style, *, centered=False, keep=False, indent=0.5):
        story.append(Paragraph(markup(value), style))
        paragraph = word.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Inches(indent)
        paragraph.paragraph_format.keep_with_next = keep
        if centered:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_word_text(paragraph, value)
        return paragraph

    for index, (kind, value) in enumerate(parsed):
        if kind == "heading":
            level, label = value
            if level == 1:
                story.extend(
                    [Spacer(1, 72), Paragraph(markup(f"**{label}**"), heading), PageBreak()]
                )
                paragraph = word.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Inches(0)
                paragraph.paragraph_format.space_before = Pt(72)
                add_word_text(paragraph, f"**{label}**")
                word.add_page_break()
            else:
                centered = level == 2 or label in {
                    "Requirements and Deliverable Traceability",
                    "Success Criteria and Acceptance Evidence",
                }
                add_paragraph(
                    f"**{label}**",
                    heading if centered else left_heading,
                    centered=centered,
                    keep=True,
                    indent=0,
                )
        elif kind == "marker":
            if value == "pagebreak":
                story.append(PageBreak())
                word.add_page_break()
            elif value == "body":
                state = "body"
                story.append(Marker("body"))
                add_paragraph(f"**{title}**", heading, centered=True, keep=True, indent=0)
            elif value == "references" or value.startswith("appendix:"):
                state = value
                story.extend([PageBreak(), Marker(value)])
                word.add_page_break()
        elif kind == "image":
            alternate, filename = value
            path = (source.parent / filename).resolve()
            if not path.is_relative_to(ROOT) or not path.is_file():
                raise ValueError(f"Invalid report image: {filename}")
            image = Image(str(path))
            factor = min(WIDTH / image.imageWidth, 320 / image.imageHeight)
            image.drawWidth, image.drawHeight = (
                image.imageWidth * factor,
                image.imageHeight * factor,
            )
            image.keepWithNext = True
            story.append(image)
            paragraph = word.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.keep_with_next = True
            picture = paragraph.add_run().add_picture(str(path), width=Pt(image.drawWidth))
            picture._inline.docPr.set("descr", alternate)
            caption_pending = False
        elif kind == "table":
            table = LongTable(
                [
                    [Paragraph(markup(f"**{v}**" if i == 0 else v), cell) for v in row]
                    for i, row in enumerate(value)
                ],
                colWidths=[Inches(w).pt for w in column_weights(value)],
                repeatRows=1,
                hAlign="LEFT",
            )
            table.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LINEABOVE", (0, 0), (-1, 0), 0.75, colors.black),
                        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
                        ("LINEBELOW", (0, -1), (-1, -1), 0.75, colors.black),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            if index + 1 < len(parsed):
                next_kind, next_value = parsed[index + 1]
                table.keepWithNext = next_kind == "paragraph" and next_value.startswith("*Note.*")
            story.append(table)
            word_table(word, value)
            caption_pending = False
        else:
            number = bool(re.fullmatch(r"\*\*(?:Table|Figure) [AB]?\d+\*\*", value))
            caption = number or caption_pending
            if state == "references":
                paragraph = add_paragraph(value, reference, indent=-0.5)
                paragraph.paragraph_format.left_indent = Inches(0.5)
            else:
                plain = state == "abstract" or caption or value.startswith("*Note.*")
                add_paragraph(
                    value,
                    caption_style if caption else no_indent if plain else body,
                    keep=caption,
                    indent=0 if plain else 0.5,
                )
            caption_pending = number

    pdf = output / "project_report.pdf"
    document = ReportDocument(pdf, title, font)
    document.build(story)
    docx = output / "project_report.docx"
    word.save(docx)
    reader = PdfReader(pdf)
    body_pages = document.markers["references"] - document.markers["body"]
    full_text = "\n".join(page.extract_text() for page in reader.pages)
    required = set(
        re.findall(
            r"\b(?:FR|TR|PR|SE|D|SC)-\d{2}\b",
            (ROOT / "docs/REQUIREMENTS.md").read_text(encoding="utf-8"),
        )
    )
    if len(reader.pages) < 7 or body_pages < 7:
        raise ValueError(
            "The report must exceed six main-body pages, "
            "excluding preliminary material and appendices"
        )
    if missing := sorted(identifier for identifier in required if identifier not in full_text):
        raise ValueError(f"Report omitted acceptance identifiers: {missing}")
    if reader.metadata.get("/Author") or word.core_properties.author:
        raise ValueError("The report must not include an author name")
    return {
        "pdf_pages": len(reader.pages),
        "main_body_pages": body_pages,
        "page_markers": document.markers,
        "requirement_identifiers_verified": len(required),
        "author_name": "",
        "font": font,
        "body_font_size": 12,
        "body_leading": 24,
        "margins_inches": 1,
        "pdf": str(pdf),
        "docx": str(docx),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "docs/project_report.md")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "docs")
    parser.add_argument("--preview-dir", type=Path)
    args = parser.parse_args()
    report = build_report(args.source, args.output_dir)
    if args.preview_dir is not None:
        import pymupdf

        args.preview_dir.mkdir(parents=True, exist_ok=True)
        with pymupdf.open(report["pdf"]) as document:
            for number, page in enumerate(document, 1):
                page.get_pixmap(matrix=pymupdf.Matrix(1, 1)).save(
                    args.preview_dir / f"page-{number:02}.png"
                )
        (args.preview_dir / "validation.json").write_text(
            json.dumps(report, indent=2) + "\n", encoding="utf-8"
        )
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
